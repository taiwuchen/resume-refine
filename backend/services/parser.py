import uuid
from pathlib import Path
from docx import Document
from models import ParsedDocument, Paragraph, ParagraphRun, PositionMapping


def _is_list_paragraph(paragraph) -> tuple[bool, int]:
    p_pr = getattr(paragraph._p, "pPr", None)
    if p_pr is None:
        return False, 0

    num_pr = getattr(p_pr, "numPr", None)
    if num_pr is not None:
        ilvl = getattr(num_pr, "ilvl", None)
        try:
            level = int(ilvl.val) if ilvl is not None and ilvl.val is not None else 0
        except (TypeError, ValueError):
            level = 0
        return True, level

    style_name = getattr(paragraph.style, "name", "") or ""
    if "list" in style_name.lower():
        return True, 0

    return False, 0


def parse_docx(file_path: Path) -> ParsedDocument:
    """Parse DOCX into full text with position→DOCX structure mapping."""
    doc = Document(file_path)
    
    full_text_parts: list[str] = []
    position_map: list[PositionMapping] = []
    paragraphs: list[Paragraph] = []
    current_pos = 0
    
    for para_idx, para in enumerate(doc.paragraphs):
        if para_idx > 0:
            full_text_parts.append("\n")
            current_pos += 1

        para_start = current_pos
        paragraph_runs: list[ParagraphRun] = []
        is_list_item, list_level = _is_list_paragraph(para)

        for run_idx, run in enumerate(para.runs):
            if not run.text:
                continue
            
            text = run.text
            start = current_pos
            end = current_pos + len(text)
            
            position_map.append(PositionMapping(
                start=start,
                end=end,
                para_idx=para_idx,
                run_idx=run_idx,
            ))

            paragraph_runs.append(ParagraphRun(
                text=text,
                start=start,
                end=end,
                bold=bool(run.bold),
                italic=bool(run.italic),
                underline=bool(run.underline),
            ))

            full_text_parts.append(text)
            current_pos = end

        paragraph_text = "".join(run.text for run in paragraph_runs)
        paragraphs.append(Paragraph(
            paragraph_id=f"p{para_idx}",
            start=para_start,
            end=current_pos,
            text=paragraph_text,
            runs=paragraph_runs,
            is_editable=bool(paragraph_text.strip()),
            is_list_item=is_list_item,
            list_level=list_level,
        ))
    
    return ParsedDocument(
        doc_id=str(uuid.uuid4()),
        full_text="".join(full_text_parts),
        position_map=position_map,
        paragraphs=paragraphs,
    )


def find_docx_locations(doc: ParsedDocument, start: int, end: int) -> list[PositionMapping]:
    """Find DOCX paragraph/run locations that overlap with the given character range."""
    locations = []
    for mapping in doc.position_map:
        if mapping.end > start and mapping.start < end:
            locations.append(mapping)
    return locations
