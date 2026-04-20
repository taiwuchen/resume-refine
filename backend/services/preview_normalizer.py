import logging
import re
from collections import Counter
from copy import deepcopy
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

logger = logging.getLogger(__name__)

MIN_HEADER_TAB_COUNT = 3
MIN_RIGHT_COLUMN_WIDTH = int(Inches(2.4))
MAX_RIGHT_COLUMN_WIDTH = int(Inches(2.85))
TARGET_RIGHT_COLUMN_RATIO = 0.4
RIGHT_VALUE_MAX_LENGTH = 44
LEFT_VALUE_MAX_LENGTH = 140
MANUAL_SPACE_SEPARATOR_RE = re.compile(r" {3,}")
WHITESPACE_RE = re.compile(r"\s+")
MONTH_PATTERN = (
    r"Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|"
    r"Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|"
    r"Nov(?:ember)?|Dec(?:ember)?"
)
YEAR_PATTERN = r"\d{4}"
LOCATION_RE = re.compile(r"^[A-Za-z .&'/()-]+,\s*[A-Z]{2,4}$")
DATE_RE = re.compile(
    rf"^(?:"
    rf"(?:{MONTH_PATTERN})\s+{YEAR_PATTERN}"
    rf"|{YEAR_PATTERN}"
    rf"|(?:{MONTH_PATTERN})(?:\s*[-/]\s*(?:{MONTH_PATTERN}))?\s+{YEAR_PATTERN}"
    rf")"
    rf"(?:\s*[-/]\s*(?:Present|Current|{YEAR_PATTERN}|(?:{MONTH_PATTERN})\s+{YEAR_PATTERN}))?$",
    re.IGNORECASE,
)
SHORT_METADATA_VALUES = {
    "remote",
    "hybrid",
    "on-site",
    "onsite",
    "contract",
    "full-time",
    "part-time",
    "internship",
}


@dataclass
class HeaderRowCandidate:
    rule_name: str
    left_children: list[object]
    right_children: list[object]
    left_text: str
    right_text: str


@dataclass
class PreviewNormalizationStats:
    counts: Counter = field(default_factory=Counter)

    def record(self, rule_name: str) -> None:
        self.counts[rule_name] += 1

    @property
    def total(self) -> int:
        return sum(self.counts.values())

    def summary(self) -> str:
        return ", ".join(
            f"{rule_name}={count}"
            for rule_name, count in sorted(self.counts.items())
        )


def normalize_preview_docx(docx_path: Path) -> int:
    doc = Document(docx_path)
    content_width = _get_content_width(doc)
    stats = PreviewNormalizationStats()

    for paragraph in list(doc.paragraphs):
        candidate = _detect_header_row_candidate(paragraph)
        if candidate is None:
            continue

        try:
            table = _insert_header_table_before(paragraph)
            _configure_header_table(table, content_width)

            left_paragraph = table.rows[0].cells[0].paragraphs[0]
            right_paragraph = table.rows[0].cells[1].paragraphs[0]

            _populate_table_paragraph(
                left_paragraph,
                paragraph,
                candidate.left_children,
                WD_PARAGRAPH_ALIGNMENT.LEFT,
            )
            _populate_table_paragraph(
                right_paragraph,
                paragraph,
                candidate.right_children,
                WD_PARAGRAPH_ALIGNMENT.RIGHT,
            )

            _remove_paragraph(paragraph)
            stats.record(candidate.rule_name)
            logger.debug(
                "Preview normalized header row via %s: left=%r right=%r",
                candidate.rule_name,
                candidate.left_text,
                candidate.right_text,
            )
        except Exception:
            logger.exception(
                "Skipping preview normalization for paragraph via %s: %r",
                candidate.rule_name,
                paragraph.text[:160],
            )

    if stats.total:
        doc.save(docx_path)
        logger.info(
            "Normalized %s preview header rows (%s)",
            stats.total,
            stats.summary(),
        )

    return stats.total


def _detect_header_row_candidate(paragraph) -> HeaderRowCandidate | None:
    if not paragraph.text.strip():
        return None

    for matcher in (
        _match_distributed_tab_header_row,
        _match_tab_stop_header_row,
        _match_manual_spacing_header_row,
    ):
        try:
            candidate = matcher(paragraph)
        except Exception:
            logger.exception(
                "Header row matcher %s failed for paragraph: %r",
                matcher.__name__,
                paragraph.text[:160],
            )
            continue

        if candidate is not None:
            return candidate

    return None


def _match_distributed_tab_header_row(paragraph) -> HeaderRowCandidate | None:
    if paragraph.alignment != WD_PARAGRAPH_ALIGNMENT.DISTRIBUTE:
        return None

    if paragraph.text.count("\t") < MIN_HEADER_TAB_COUNT:
        return None

    split_children = _split_tab_aligned_paragraph(paragraph)
    if split_children is None:
        return None

    return _build_candidate("distributed_tabs", *split_children)


def _match_tab_stop_header_row(paragraph) -> HeaderRowCandidate | None:
    if paragraph.text.count("\t") < 1:
        return None

    tab_stops = list(paragraph.paragraph_format.tab_stops)
    if not tab_stops:
        return None

    if not any(
        getattr(tab_stop, "alignment", None) == WD_TAB_ALIGNMENT.RIGHT
        for tab_stop in tab_stops
    ):
        return None

    split_children = _split_tab_aligned_paragraph(paragraph)
    if split_children is None:
        return None

    return _build_candidate("tab_stops", *split_children)


def _match_manual_spacing_header_row(paragraph) -> HeaderRowCandidate | None:
    if "\t" in paragraph.text or not paragraph.text.strip():
        return None

    split_offsets = _find_manual_spacing_split(paragraph.text)
    if split_offsets is None:
        return None

    left_end, right_start = split_offsets
    split_children = _split_children_by_text_offsets(paragraph, left_end, right_start)
    if split_children is None:
        return None

    return _build_candidate("manual_spacing", *split_children, strict=True)


def _find_manual_spacing_split(text: str) -> tuple[int, int] | None:
    matches = list(MANUAL_SPACE_SEPARATOR_RE.finditer(text))
    if not matches:
        return None

    for match in reversed(matches):
        left_text = text[: match.start()].rstrip()
        right_text = text[match.end() :].lstrip()
        if _is_likely_header_split(left_text, right_text, strict=True):
            return match.start(), match.end()

    return None


def _build_candidate(
    rule_name: str,
    left_children: list[object],
    right_children: list[object],
    *,
    strict: bool = False,
) -> HeaderRowCandidate | None:
    if not left_children or not right_children:
        return None

    _trim_child_edges(left_children)
    _trim_child_edges(right_children)

    left_text = _normalize_inline_text(_extract_children_text(left_children))
    right_text = _normalize_inline_text(_extract_children_text(right_children))

    if not _is_likely_header_split(left_text, right_text, strict=strict):
        return None

    return HeaderRowCandidate(
        rule_name=rule_name,
        left_children=left_children,
        right_children=right_children,
        left_text=left_text,
        right_text=right_text,
    )


def _is_likely_header_split(left_text: str, right_text: str, *, strict: bool = False) -> bool:
    left_text = _normalize_inline_text(left_text)
    right_text = _normalize_inline_text(right_text)

    if not left_text or not right_text:
        return False

    if len(right_text) > RIGHT_VALUE_MAX_LENGTH or len(left_text) > LEFT_VALUE_MAX_LENGTH:
        return False

    if right_text.endswith((".", ";", ":")) or left_text.endswith(";"):
        return False

    if not _is_likely_right_value(right_text):
        return False

    left_word_count = len(left_text.split())
    if strict and left_word_count > 10:
        return False

    if left_word_count > 18:
        return False

    if strict and len(left_text) <= len(right_text):
        return False

    if any(marker in left_text for marker in ("•", "▪", "◦")):
        return False

    return True


def _is_likely_right_value(text: str) -> bool:
    normalized = _normalize_inline_text(text)
    lowered = normalized.lower()

    if lowered in SHORT_METADATA_VALUES:
        return True

    if DATE_RE.match(normalized):
        return True

    if LOCATION_RE.match(normalized):
        return True

    if any(token in lowered for token in ("present", "current")):
        return True

    return bool(re.search(YEAR_PATTERN, normalized))


def _normalize_inline_text(text: str) -> str:
    return WHITESPACE_RE.sub(" ", text.replace("\t", " ")).strip()


def _get_content_width(doc: Document) -> int:
    section = doc.sections[0]
    return int(section.page_width - section.left_margin - section.right_margin)


def _split_tab_aligned_paragraph(paragraph) -> tuple[list[object], list[object]] | None:
    left_children: list[object] = []
    right_children: list[object] = []
    seen_separator_tabs = False

    for child in paragraph._p:
        if child.tag == qn("w:pPr"):
            continue

        has_tab = _has_tab(child)
        raw_text = _extract_text(child)
        has_text = bool(raw_text.strip())
        has_whitespace_text = bool(raw_text) and not has_text

        if has_tab and has_text:
            if not left_children:
                return None

            seen_separator_tabs = True
            right_child = _clone_child_without_tabs(child)
            if _normalize_inline_text(_extract_text(right_child)):
                right_children.append(right_child)
            continue

        if not seen_separator_tabs:
            if has_text:
                left_children.append(deepcopy(child))
                continue

            if has_whitespace_text and left_children:
                left_children.append(deepcopy(child))
                continue

            if has_tab and left_children:
                seen_separator_tabs = True
            continue

        if has_tab and not has_text:
            continue

        if has_text:
            right_children.append(deepcopy(child))
            continue

        if has_whitespace_text and right_children:
            right_children.append(deepcopy(child))

    if not left_children or not right_children:
        return None

    return left_children, right_children


def _split_children_by_text_offsets(
    paragraph,
    left_end: int,
    right_start: int,
) -> tuple[list[object], list[object]] | None:
    left_children: list[object] = []
    right_children: list[object] = []
    cursor = 0

    for child in paragraph._p:
        if child.tag == qn("w:pPr"):
            continue

        child_text = _extract_text(child)
        child_length = len(child_text)
        child_start = cursor
        child_end = cursor + child_length

        if child_length == 0:
            continue

        if child_start < left_end and child_end > 0:
            left_child = _clone_child_text_range(
                child,
                max(0, 0 - child_start),
                min(child_length, left_end - child_start),
            )
            if left_child is not None:
                left_children.append(left_child)

        if child_start < child_end and child_end > right_start:
            right_child = _clone_child_text_range(
                child,
                max(0, right_start - child_start),
                child_length,
            )
            if right_child is not None:
                right_children.append(right_child)

        cursor = child_end

    if not left_children or not right_children:
        return None

    return left_children, right_children


def _clone_child_text_range(child, start: int, end: int):
    if end <= start:
        return None

    child_text = _extract_text(child)
    if not child_text:
        return None

    if start <= 0 and end >= len(child_text):
        cloned_child = deepcopy(child)
        _preserve_text_spacing(cloned_child)
        return cloned_child

    cloned_child = deepcopy(child)
    cursor = 0

    for text_node in list(cloned_child.iter(qn("w:t"))):
        original_text = text_node.text or ""
        node_start = cursor
        node_end = cursor + len(original_text)
        overlap_start = max(start, node_start)
        overlap_end = min(end, node_end)

        if overlap_start >= overlap_end:
            text_node.text = ""
        else:
            local_start = overlap_start - node_start
            local_end = overlap_end - node_start
            text_node.text = original_text[local_start:local_end]

        cursor = node_end

    _strip_empty_content(cloned_child)
    if not _normalize_inline_text(_extract_text(cloned_child)):
        return None

    _preserve_text_spacing(cloned_child)
    return cloned_child


def _strip_empty_content(element) -> None:
    for descendant in reversed(list(element.iter())):
        if descendant is element:
            continue

        if descendant.tag == qn("w:t") and not (descendant.text or ""):
            parent = descendant.getparent()
            if parent is not None:
                parent.remove(descendant)
            continue

        if descendant.tag in {qn("w:r"), qn("w:hyperlink")}:
            if _extract_text(descendant):
                continue
            parent = descendant.getparent()
            if parent is not None:
                parent.remove(descendant)


def _extract_text(element) -> str:
    return "".join(node.text or "" for node in element.iter(qn("w:t")))


def _extract_children_text(children: list[object]) -> str:
    return "".join(_extract_text(child) for child in children)


def _has_tab(element) -> bool:
    return any(True for _ in element.iter(qn("w:tab")))


def _trim_child_edges(children: list[object]) -> None:
    if not children:
        return

    first_text_nodes = list(children[0].iter(qn("w:t")))
    if first_text_nodes:
        first_text_nodes[0].text = (first_text_nodes[0].text or "").lstrip()

    last_text_nodes = list(children[-1].iter(qn("w:t")))
    if last_text_nodes:
        last_text_nodes[-1].text = (last_text_nodes[-1].text or "").rstrip()

    for child in children:
        _preserve_text_spacing(child)


def _clone_child_without_tabs(child):
    cloned_child = deepcopy(child)

    for tab in list(cloned_child.iter(qn("w:tab"))):
        tab.getparent().remove(tab)

    _preserve_text_spacing(cloned_child)
    return cloned_child


def _preserve_text_spacing(element) -> None:
    for text_node in element.iter(qn("w:t")):
        text = text_node.text or ""
        if text.startswith(" ") or text.endswith(" ") or "  " in text:
            text_node.set(qn("xml:space"), "preserve")


def _insert_header_table_before(paragraph):
    document = paragraph.part.document
    table = document.add_table(rows=1, cols=2)
    paragraph._p.addprevious(table._tbl)
    return table


def _configure_header_table(table, content_width: int) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    _set_table_borders_none(table)
    _set_table_layout_fixed(table)

    right_width = int(content_width * TARGET_RIGHT_COLUMN_RATIO)
    right_width = max(MIN_RIGHT_COLUMN_WIDTH, min(MAX_RIGHT_COLUMN_WIDTH, right_width))
    left_width = max(0, content_width - right_width)

    for cell, width in zip(table.rows[0].cells, (left_width, right_width)):
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _set_cell_width(cell, width)
        _set_cell_margins(cell, top=0, bottom=0, left=0, right=0)


def _populate_table_paragraph(target_paragraph, source_paragraph, children: list[object], alignment) -> None:
    target_paragraph.clear()
    target_paragraph.style = source_paragraph.style
    _copy_paragraph_format(source_paragraph, target_paragraph)
    target_paragraph.alignment = alignment

    for child in children:
        target_paragraph._p.append(deepcopy(child))


def _copy_paragraph_format(source_paragraph, target_paragraph) -> None:
    source_format = source_paragraph.paragraph_format
    target_format = target_paragraph.paragraph_format
    target_format.left_indent = source_format.left_indent
    target_format.right_indent = source_format.right_indent
    target_format.first_line_indent = source_format.first_line_indent
    target_format.keep_together = source_format.keep_together
    target_format.keep_with_next = source_format.keep_with_next
    target_format.page_break_before = source_format.page_break_before
    target_format.widow_control = source_format.widow_control
    target_format.space_before = source_format.space_before
    target_format.space_after = source_format.space_after
    target_format.line_spacing = source_format.line_spacing
    target_format.line_spacing_rule = source_format.line_spacing_rule


def _set_table_borders_none(table) -> None:
    table_properties = table._tbl.tblPr
    borders = table_properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge_element = borders.find(qn(f"w:{edge}"))
        if edge_element is None:
            edge_element = OxmlElement(f"w:{edge}")
            borders.append(edge_element)
        edge_element.set(qn("w:val"), "nil")


def _set_table_layout_fixed(table) -> None:
    table_properties = table._tbl.tblPr
    layout = table_properties.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_properties.append(layout)
    layout.set(qn("w:type"), "fixed")


def _set_cell_width(cell, width: int) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    width_element = cell_properties.first_child_found_in("w:tcW")
    if width_element is None:
        width_element = OxmlElement("w:tcW")
        cell_properties.append(width_element)
    width_element.set(qn("w:type"), "dxa")
    width_element.set(qn("w:w"), str(int(width / 635)))


def _set_cell_margins(cell, *, top: int, bottom: int, left: int, right: int) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    margins = cell_properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        cell_properties.append(margins)

    for side, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        margin = margins.find(qn(f"w:{side}"))
        if margin is None:
            margin = OxmlElement(f"w:{side}")
            margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def _remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
