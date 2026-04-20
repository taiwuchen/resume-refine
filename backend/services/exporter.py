from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from models.document import Change, ParsedDocument
from services.docx.xml_utils import (
    clone_run_properties,
    extract_text,
    is_hyperlink,
    is_run,
)


@dataclass(frozen=True)
class HyperlinkSource:
    text: str
    element: object


def apply_changes_to_docx(
    original_path: Path,
    doc: ParsedDocument,
    changes: list[Change],
    output_path: Path,
) -> Path:
    """Apply text changes to original DOCX while preserving formatting."""
    docx = Document(original_path)

    paragraph_map = {
        paragraph.paragraph_id: (index, paragraph)
        for index, paragraph in enumerate(doc.paragraphs)
    }
    seen_paragraph_ids: set[str] = set()

    for change in changes:
        paragraph_entry = paragraph_map.get(change.paragraph_id)
        if paragraph_entry is None:
            raise ValueError(f"Unknown paragraph: {change.paragraph_id}")

        paragraph_index, paragraph = paragraph_entry
        if not paragraph.is_editable:
            raise ValueError(f"Paragraph is not editable: {change.paragraph_id}")

        if change.paragraph_id in seen_paragraph_ids:
            raise ValueError(f"Duplicate change for paragraph: {change.paragraph_id}")

        if change.start != paragraph.start or change.end != paragraph.end:
            raise ValueError("Changes must target a full paragraph")

        if change.original != paragraph.text:
            raise ValueError("Change original text does not match the paragraph")

        apply_single_change(docx, paragraph_index, change.replacement)
        seen_paragraph_ids.add(change.paragraph_id)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    docx.save(output_path)
    return output_path
def _append_text_children(run_element, text: str) -> None:
    parts = text.split("\t")

    for part_index, part in enumerate(parts):
        line_parts = part.split("\n")
        for line_index, line in enumerate(line_parts):
            if line:
                text_element = OxmlElement("w:t")
                if line != line.strip() or "  " in line:
                    text_element.set(qn("xml:space"), "preserve")
                text_element.text = line
                run_element.append(text_element)

            if line_index < len(line_parts) - 1:
                run_element.append(OxmlElement("w:br"))

        if part_index < len(parts) - 1:
            run_element.append(OxmlElement("w:tab"))

    if not text:
        text_element = OxmlElement("w:t")
        text_element.text = ""
        run_element.append(text_element)


def _build_plain_run(text: str, style_source) -> object:
    run_element = OxmlElement("w:r")
    if style_source is not None:
        clone_run_properties(style_source, run_element)
    _append_text_children(run_element, text)
    return run_element


def _get_representative_run(paragraph) -> object | None:
    for child in paragraph._p:
        if is_run(child) and extract_text(child):
            return child
        if is_hyperlink(child):
            for hyperlink_child in child:
                if is_run(hyperlink_child) and extract_text(hyperlink_child):
                    return hyperlink_child
    return None


def _get_plain_style_source(paragraph) -> object | None:
    for child in paragraph._p:
        if is_run(child) and extract_text(child):
            return child
    return _get_representative_run(paragraph)


def _get_hyperlink_sources(paragraph) -> list[HyperlinkSource]:
    sources: list[HyperlinkSource] = []

    for child in paragraph._p:
        if not is_hyperlink(child):
            continue

        text = extract_text(child)
        if text:
            sources.append(HyperlinkSource(text=text, element=child))

    return sources


def _build_paragraph_children(paragraph, replacement: str) -> list[object]:
    plain_style_source = _get_plain_style_source(paragraph)
    hyperlink_sources = _get_hyperlink_sources(paragraph)

    if not hyperlink_sources:
        return [_build_plain_run(replacement, plain_style_source)]

    children: list[object] = []
    cursor = 0
    preserved_hyperlinks = 0

    for source in hyperlink_sources:
        match_index = replacement.find(source.text, cursor)
        if match_index == -1:
            continue

        if match_index > cursor:
            children.append(_build_plain_run(replacement[cursor:match_index], plain_style_source))

        children.append(deepcopy(source.element))
        cursor = match_index + len(source.text)
        preserved_hyperlinks += 1

    if cursor < len(replacement):
        children.append(_build_plain_run(replacement[cursor:], plain_style_source))

    if preserved_hyperlinks == 0:
        return [_build_plain_run(replacement, plain_style_source)]

    return children or [_build_plain_run("", plain_style_source)]


def apply_single_change(
    docx: Document,
    paragraph_index: int,
    replacement: str,
) -> None:
    """Apply a single full-paragraph text change to the DOCX document."""
    if paragraph_index >= len(docx.paragraphs):
        raise ValueError(f"Paragraph index out of range: {paragraph_index}")

    paragraph = docx.paragraphs[paragraph_index]
    paragraph_children = _build_paragraph_children(paragraph, replacement)

    paragraph.clear()
    for child in paragraph_children:
        paragraph._p.append(child)
