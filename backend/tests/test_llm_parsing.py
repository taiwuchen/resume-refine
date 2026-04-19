import tempfile
import unittest
from pathlib import Path

from docx import Document

from services.analyzer import parse_analyze_response
from services.chat import parse_chat_response
from services.parser import parse_docx


def build_parsed_document():
    temp_dir = tempfile.TemporaryDirectory()
    file_path = Path(temp_dir.name) / "resume.docx"
    doc = Document()
    doc.add_paragraph("Built backend systems")
    doc.add_paragraph("")
    doc.add_paragraph("Scaled repeated workflows")
    doc.save(file_path)
    parsed = parse_docx(file_path)
    return temp_dir, parsed


class AnalyzerParsingTests(unittest.TestCase):
    def test_parse_analyze_response_validates_ids_and_deduplicates(self) -> None:
        temp_dir, parsed = build_parsed_document()
        self.addCleanup(temp_dir.cleanup)

        raw = """
        {
            "suggestions": [
                {
                    "paragraph_id": "p0",
                    "alternatives": ["Alt one", "Alt two", "Alt three"]
                },
                {
                    "paragraph_id": "p0",
                    "alternatives": ["Duplicate one", "Duplicate two", "Duplicate three"]
                },
                {
                    "paragraph_id": "p1",
                    "alternatives": ["Blank one", "Blank two", "Blank three"]
                },
                {
                    "paragraph_id": "p99",
                    "alternatives": ["Missing one", "Missing two", "Missing three"]
                },
                {
                    "paragraph_id": "p2",
                    "alternatives": ["Only one"]
                }
            ]
        }
        """

        suggestions = parse_analyze_response(raw, parsed)

        self.assertEqual(1, len(suggestions))
        self.assertEqual("p0", suggestions[0].paragraph_id)
        self.assertEqual("Built backend systems", suggestions[0].original_text)
        self.assertEqual(["Alt one", "Alt two", "Alt three"], suggestions[0].alternatives)

    def test_parse_analyze_response_returns_empty_on_malformed_json(self) -> None:
        temp_dir, parsed = build_parsed_document()
        self.addCleanup(temp_dir.cleanup)

        suggestions = parse_analyze_response("{not json", parsed)

        self.assertEqual([], suggestions)


class ChatParsingTests(unittest.TestCase):
    def test_parse_chat_response_resolves_valid_paragraph_edits(self) -> None:
        temp_dir, parsed = build_parsed_document()
        self.addCleanup(temp_dir.cleanup)

        raw = """
        {
            "message": "Here are better options.",
            "edits": [
                {
                    "paragraph_id": "p0",
                    "new_text": "Built reliable backend systems at scale",
                    "explanation": "Adds impact."
                },
                {
                    "paragraph_id": "p1",
                    "new_text": "Should be ignored",
                    "explanation": "Blank paragraph"
                },
                {
                    "paragraph_id": "p99",
                    "new_text": "Missing paragraph",
                    "explanation": "Unknown"
                },
                {
                    "paragraph_id": "p0",
                    "new_text": "Duplicate edit",
                    "explanation": "Duplicate"
                }
            ]
        }
        """

        response = parse_chat_response(raw, parsed)

        self.assertEqual("Here are better options.", response.message)
        self.assertEqual(1, len(response.edits))
        self.assertEqual("p0", response.edits[0].paragraph_id)
        self.assertEqual("Built backend systems", response.edits[0].original_text)
        self.assertEqual("Built reliable backend systems at scale", response.edits[0].new_text)
        self.assertEqual(parsed.paragraphs[0].start, response.edits[0].start)
        self.assertEqual(parsed.paragraphs[0].end, response.edits[0].end)

    def test_parse_chat_response_returns_raw_message_on_invalid_json(self) -> None:
        temp_dir, parsed = build_parsed_document()
        self.addCleanup(temp_dir.cleanup)

        response = parse_chat_response("not valid json", parsed)

        self.assertEqual("not valid json", response.message)
        self.assertEqual([], response.edits)


if __name__ == "__main__":
    unittest.main()
