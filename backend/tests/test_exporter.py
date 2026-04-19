import tempfile
import unittest
from pathlib import Path

from docx import Document

from models import Change
from services.exporter import apply_changes_to_docx
from services.parser import parse_docx


def create_export_fixture(temp_dir: str) -> tuple[Path, object]:
    file_path = Path(temp_dir) / "resume.docx"
    doc = Document()

    first_paragraph = doc.add_paragraph()
    first_paragraph.add_run("Led ")
    highlighted_run = first_paragraph.add_run("projects")
    highlighted_run.bold = True

    second_paragraph = doc.add_paragraph(style="List Bullet")
    second_paragraph.add_run("Shipped customer features")

    doc.save(file_path)
    return file_path, parse_docx(file_path)


class ExporterTests(unittest.TestCase):
    def test_apply_changes_replaces_only_targeted_paragraphs(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            original_path, parsed = create_export_fixture(temp_dir)
            output_path = Path(temp_dir) / "refined.docx"

            apply_changes_to_docx(
                original_path,
                parsed,
                [
                    Change(
                        paragraph_id="p0",
                        start=parsed.paragraphs[0].start,
                        end=parsed.paragraphs[0].end,
                        original=parsed.paragraphs[0].text,
                        replacement="Directed strategic programs",
                    ),
                    Change(
                        paragraph_id="p1",
                        start=parsed.paragraphs[1].start,
                        end=parsed.paragraphs[1].end,
                        original=parsed.paragraphs[1].text,
                        replacement="Delivered customer-facing features",
                    ),
                ],
                output_path,
            )

            exported = Document(output_path)

        self.assertEqual("Directed strategic programs", exported.paragraphs[0].text)
        self.assertEqual("Delivered customer-facing features", exported.paragraphs[1].text)
        self.assertEqual("Directed strategic programs", exported.paragraphs[0].runs[0].text)
        self.assertEqual("", exported.paragraphs[0].runs[1].text)

    def test_apply_changes_rejects_partial_or_duplicate_changes(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            original_path, parsed = create_export_fixture(temp_dir)
            output_path = Path(temp_dir) / "refined.docx"

            with self.assertRaisesRegex(ValueError, "full paragraph"):
                apply_changes_to_docx(
                    original_path,
                    parsed,
                    [
                        Change(
                            paragraph_id="p0",
                            start=parsed.paragraphs[0].start + 1,
                            end=parsed.paragraphs[0].end,
                            original=parsed.paragraphs[0].text,
                            replacement="Invalid partial update",
                        )
                    ],
                    output_path,
                )

            with self.assertRaisesRegex(ValueError, "Duplicate change"):
                apply_changes_to_docx(
                    original_path,
                    parsed,
                    [
                        Change(
                            paragraph_id="p0",
                            start=parsed.paragraphs[0].start,
                            end=parsed.paragraphs[0].end,
                            original=parsed.paragraphs[0].text,
                            replacement="First update",
                        ),
                        Change(
                            paragraph_id="p0",
                            start=parsed.paragraphs[0].start,
                            end=parsed.paragraphs[0].end,
                            original=parsed.paragraphs[0].text,
                            replacement="Second update",
                        ),
                    ],
                    output_path,
                )


if __name__ == "__main__":
    unittest.main()
