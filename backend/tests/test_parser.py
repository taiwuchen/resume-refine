import tempfile
import unittest
from pathlib import Path

from docx import Document

from services.parser import parse_docx


class ParserTests(unittest.TestCase):
    def test_parse_docx_preserves_paragraph_anchors_and_runs(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            file_path = Path(temp_dir) / "resume.docx"
            doc = Document()

            first_paragraph = doc.add_paragraph()
            first_paragraph.add_run("Led ")
            highlighted_run = first_paragraph.add_run("projects")
            highlighted_run.bold = True

            doc.add_paragraph("")

            bullet_one = doc.add_paragraph(style="List Bullet")
            bullet_one.add_run("Repeated bullet")

            bullet_two = doc.add_paragraph(style="List Bullet")
            bullet_two.add_run("Repeated bullet")

            doc.save(file_path)

            parsed = parse_docx(file_path)

        self.assertEqual(4, len(parsed.paragraphs))
        self.assertEqual("Led projects", parsed.paragraphs[0].text)
        self.assertTrue(parsed.paragraphs[0].is_editable)
        self.assertEqual(2, len(parsed.paragraphs[0].runs))
        self.assertFalse(parsed.paragraphs[0].runs[0].bold)
        self.assertTrue(parsed.paragraphs[0].runs[1].bold)

        self.assertEqual("", parsed.paragraphs[1].text)
        self.assertFalse(parsed.paragraphs[1].is_editable)

        self.assertTrue(parsed.paragraphs[2].is_list_item)
        self.assertTrue(parsed.paragraphs[3].is_list_item)
        self.assertEqual("Repeated bullet", parsed.paragraphs[2].text)
        self.assertEqual("Repeated bullet", parsed.paragraphs[3].text)
        self.assertNotEqual(
            parsed.paragraphs[2].paragraph_id,
            parsed.paragraphs[3].paragraph_id,
        )

        self.assertIn("Led projects\n\nRepeated bullet\nRepeated bullet", parsed.full_text)


if __name__ == "__main__":
    unittest.main()
