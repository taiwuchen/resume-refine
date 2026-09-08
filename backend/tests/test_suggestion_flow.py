import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from main import app
from models.document import Change
from services.ai.analyzer import InvalidAnalysisResponse, parse_analyze_response
from services.exporter import apply_changes_to_docx
from services.parser import parse_docx
from services.resume_state import apply_changes_to_parsed_document


class SuggestionFlowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'resume.docx'
        document = Document()
        document.add_paragraph('Built reporting tools for the operations team.')
        document.add_paragraph('Worked with stakeholders to improve reporting.')
        document.save(self.source)
        self.doc = parse_docx(self.source)
        self.client = TestClient(app)

    def change(self, **overrides):
        paragraph = self.doc.paragraphs[0]
        fields = dict(paragraph_id=paragraph.paragraph_id, start=paragraph.start, end=paragraph.end,
                      original=paragraph.text, replacement='Built operations reporting tools.')
        return Change(**(fields | overrides))

    def test_analysis_rejects_malformed_results_and_accepts_explicit_empty(self):
        for response in ['not JSON', '{}', '[]', '{"suggestions":[{}]}']:
            with self.subTest(response=response), self.assertRaises(InvalidAnalysisResponse):
                parse_analyze_response(response, self.doc)
        self.assertEqual(parse_analyze_response('{"suggestions":[]}', self.doc), [])

    def test_suggestions_use_current_text_and_original_export_anchors(self):
        effective = apply_changes_to_parsed_document(self.doc, [self.change()])
        result = parse_analyze_response(json.dumps({'suggestions': [{
            'paragraph_id': 'p1', 'reason': 'Clarify collaboration.',
            'alternatives': ['Partnered with stakeholders on reporting.', 'Improved reporting with stakeholders.', 'Collaborated on reporting improvements.'],
        }, {
            'paragraph_id': 'p0', 'reason': 'Make the wording concise.',
            'alternatives': ['Created operations reporting tools.', 'Developed tools for operations reporting.', 'Delivered reporting tools for operations.'],
        }]}), effective)
        self.assertEqual(result[0].paragraph_id, 'p0')
        self.assertEqual(result[0].original_text, self.change().replacement)
        self.assertEqual(result[0].end, self.doc.paragraphs[0].end)
        output = self.root / 'result.docx'
        apply_changes_to_docx(self.source, self.doc, [self.change(replacement=result[0].alternatives[0])], output)
        self.assertEqual(Document(output).paragraphs[0].text, result[0].alternatives[0])
        apply_changes_to_docx(self.source, self.doc, [], output)
        self.assertEqual(Document(output).paragraphs[0].text, self.doc.paragraphs[0].text)

    def test_analysis_and_export_reject_the_same_invalid_changes(self):
        for changes in [[self.change(original='stale')], [self.change(paragraph_id='missing')],
                        [self.change(start=1)], [self.change(), self.change()]]:
            with self.subTest(changes=changes):
                with self.assertRaises(ValueError):
                    apply_changes_to_parsed_document(self.doc, changes)
                with self.assertRaises(ValueError):
                    apply_changes_to_docx(self.source, self.doc, changes, self.root / 'invalid.docx')

    def test_uploads_with_identical_names_are_isolated(self):
        stored = {}

        def store(doc, path, token):
            stored[doc.doc_id] = path

        with patch('routers.upload.UPLOAD_DIR', self.root), \
             patch('routers.upload.document_repository.store_document', side_effect=store):
            first = self.client.post('/api/upload', files={'file': ('resume.docx', self.source.read_bytes())})
            document = Document()
            document.add_paragraph('A different resume.')
            data = io.BytesIO()
            document.save(data)
            second = self.client.post('/api/upload', files={'file': ('resume.docx', data.getvalue())})
        self.assertEqual(first.status_code, 200)
        self.assertEqual(second.status_code, 200)
        first_path = stored[first.json()['document']['doc_id']]
        second_path = stored[second.json()['document']['doc_id']]
        self.assertNotEqual(first_path, second_path)
        self.assertEqual(Document(first_path).paragraphs[0].text, self.doc.paragraphs[0].text)
        self.assertEqual(Document(second_path).paragraphs[0].text, 'A different resume.')

    def test_analysis_failure_is_not_a_successful_empty_result(self):
        with patch('routers.dependencies.document_repository.get_document', return_value=self.doc), \
             patch('services.ai.analyzer.call_llm', return_value='invalid'):
            response = self.client.post(
                '/api/analyze',
                json={'doc_id': self.doc.doc_id, 'job_description': 'Reporting engineer'},
                headers={'X-Document-Token': 'token', 'X-OpenRouter-Key': 'sk-test'},
            )
        self.assertEqual(response.status_code, 502)
        self.assertIn('try again', response.json()['detail'])

    def test_removed_endpoints_are_unavailable(self):
        for endpoint in ['/api/chat', '/api/suggest']:
            self.assertEqual(self.client.post(endpoint, json={}).status_code, 404)


class UploadLimitTests(unittest.TestCase):
    """Uploads are anonymous, so every limit here is what stands between one
    request and an out-of-memory instance."""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.client = TestClient(app)

    def _docx_bytes(self, text='Resume line.'):
        document = Document()
        document.add_paragraph(text)
        data = io.BytesIO()
        document.save(data)
        return data.getvalue()

    def _bomb_bytes(self):
        """A small archive whose document.xml expands enormously."""
        import zipfile

        base = io.BytesIO(self._docx_bytes())
        with zipfile.ZipFile(base) as source:
            items = {name: source.read(name) for name in source.namelist()}

        body = b'<w:p><w:r><w:t>' + b'A' * 200 + b'</w:t></w:r></w:p>'
        items['word/document.xml'] = (
            b'<?xml version="1.0"?><w:document '
            b'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
            + body * 20_000
            + b'</w:body></w:document>'
        )

        out = io.BytesIO()
        with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as target:
            for name, payload in items.items():
                target.writestr(name, payload)
        return out.getvalue()

    def test_upload_over_the_size_limit_is_rejected(self):
        with patch('routers.upload.MAX_UPLOAD_BYTES', 1024), patch('routers.upload.UPLOAD_DIR', self.root):
            response = self.client.post(
                '/api/upload',
                files={'file': ('resume.docx', b'x' * 8192)},
            )
        self.assertEqual(response.status_code, 413)
        self.assertEqual(list(self.root.iterdir()), [], 'oversized upload must not be left on disk')

    def test_decompression_bomb_is_rejected_before_parsing(self):
        bomb = self._bomb_bytes()
        self.assertLess(len(bomb), 200_000, 'bomb should be small on the wire')

        with patch('services.docx_safety.MAX_EXPANDED_BYTES', 1_000_000), \
             patch('routers.upload.UPLOAD_DIR', self.root):
            response = self.client.post('/api/upload', files={'file': ('resume.docx', bomb)})

        self.assertEqual(response.status_code, 400)
        self.assertEqual(list(self.root.iterdir()), [], 'rejected upload must not be left on disk')

    def test_non_docx_archive_is_rejected(self):
        with patch('routers.upload.UPLOAD_DIR', self.root):
            response = self.client.post('/api/upload', files={'file': ('resume.docx', b'not a zip at all')})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(list(self.root.iterdir()), [])


class DocumentAccessTests(unittest.TestCase):
    """A doc_id alone must not grant access to someone else's resume."""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.source = self.root / 'resume.docx'
        document = Document()
        document.add_paragraph('Built reporting tools for the operations team.')
        document.save(self.source)
        self.client = TestClient(app)

    def _upload(self):
        with patch('routers.upload.UPLOAD_DIR', self.root):
            response = self.client.post(
                '/api/upload', files={'file': ('resume.docx', self.source.read_bytes())}
            )
        self.assertEqual(response.status_code, 200)
        body = response.json()
        return body['document']['doc_id'], body['access_token']

    def test_upload_issues_an_access_token(self):
        doc_id, token = self._upload()
        self.assertTrue(token)
        self.assertNotEqual(token, doc_id)

    def test_document_requires_the_matching_token(self):
        doc_id, token = self._upload()

        self.assertEqual(self.client.get(f'/api/document/{doc_id}').status_code, 404)
        self.assertEqual(
            self.client.get(f'/api/document/{doc_id}', headers={'X-Document-Token': 'wrong'}).status_code,
            404,
        )
        self.assertEqual(
            self.client.get(f'/api/document/{doc_id}', headers={'X-Document-Token': token}).status_code,
            200,
        )

    def test_export_requires_the_matching_token(self):
        doc_id, token = self._upload()
        payload = {'doc_id': doc_id, 'changes': []}

        self.assertEqual(
            self.client.post('/api/export', json=payload, headers={'X-Document-Token': 'wrong'}).status_code,
            404,
        )
        self.assertEqual(
            self.client.post('/api/export', json=payload, headers={'X-Document-Token': token}).status_code,
            200,
        )


class ApiKeyTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)
        document = Document()
        document.add_paragraph('Built reporting tools for the operations team.')
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        source = Path(self.temp.name) / 'resume.docx'
        document.save(source)
        self.doc = parse_docx(source)

    def test_analysis_without_a_caller_key_is_refused(self):
        """The server key must not fund anonymous callers."""
        with patch('routers.dependencies.document_repository.get_document', return_value=self.doc), \
             patch('services.ai.llm.ALLOW_SERVER_API_KEY', False), \
             patch('services.ai.llm.OPENROUTER_API_KEY', 'server-key'):
            response = self.client.post(
                '/api/analyze',
                json={'doc_id': self.doc.doc_id, 'job_description': 'Reporting engineer'},
                headers={'X-Document-Token': 'token'},
            )
        self.assertEqual(response.status_code, 401)
        self.assertIn('API key', response.json()['detail'])

    def test_caller_key_is_used_and_not_echoed(self):
        seen = {}

        def fake_call(messages, api_key):
            seen['key'] = api_key
            return '{"suggestions":[]}'

        with patch('routers.dependencies.document_repository.get_document', return_value=self.doc), \
             patch('services.ai.analyzer.call_llm', side_effect=fake_call):
            response = self.client.post(
                '/api/analyze',
                json={'doc_id': self.doc.doc_id, 'job_description': 'Reporting engineer'},
                headers={'X-Document-Token': 'token', 'X-OpenRouter-Key': 'sk-caller'},
            )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(seen['key'], 'sk-caller')
        self.assertNotIn('sk-caller', response.text)

    def test_job_description_over_the_limit_is_rejected(self):
        with patch('routers.dependencies.document_repository.get_document', return_value=self.doc):
            response = self.client.post(
                '/api/analyze',
                json={'doc_id': self.doc.doc_id, 'job_description': 'x' * 200_000},
                headers={'X-Document-Token': 'token', 'X-OpenRouter-Key': 'sk-caller'},
            )
        self.assertEqual(response.status_code, 422)


if __name__ == '__main__':
    unittest.main()
