from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph as DocxParagraph
from docx2pdf import convert

WORD_BRIDGE_DIR = Path.home() / ".resume-refine" / "word-bridge"
WORD_BRIDGE_DOCX = WORD_BRIDGE_DIR / "bridge.docx"
WORD_BRIDGE_PDF = WORD_BRIDGE_DIR / "bridge.pdf"


def _replace_paragraph_text(paragraph: DocxParagraph, new_text: str) -> None:
    """Replace text while preserving paragraph formatting."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(new_text)
        return

    runs[0].text = new_text
    for run in runs[1:]:
        run.text = ""


def apply_changes(doc: Document, optimized: dict[int, str]) -> Document:
    for idx, new_text in optimized.items():
        if idx < 0 or idx >= len(doc.paragraphs):
            raise IndexError(f"Paragraph index {idx} out of range")

        paragraph = doc.paragraphs[idx]
        _replace_paragraph_text(paragraph, new_text)

    return doc


def save_docx(doc: Document, output_path: str) -> str:
    path = Path(output_path).expanduser()
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return str(path)


def export_pdf(docx_path: str | Path, output_path: str) -> str:
    """
    Convert the provided DOCX file to PDF via Microsoft Word.

    Word (especially on macOS) frequently prompts for file-access permissions
    whenever it sees a new path. We always convert through a stable "bridge"
    file so the user only has to grant permission once.
    """
    source = Path(docx_path).expanduser().resolve()
    if not source.exists():
        raise FileNotFoundError(f"DOCX not found for PDF export: {source}")

    output = Path(output_path).expanduser().resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    WORD_BRIDGE_DIR.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, WORD_BRIDGE_DOCX)

    try:
        convert(str(WORD_BRIDGE_DOCX), str(WORD_BRIDGE_PDF))
    except SystemExit as exc:  # docx2pdf uses sys.exit on failure
        raise RuntimeError(
            "docx2pdf failed while asking Microsoft Word to export the PDF. "
            "If macOS shows a 'Grant File Access' dialog for the bridge file "
            f"({WORD_BRIDGE_DOCX}), approve it once and rerun."
        ) from exc

    shutil.copy2(WORD_BRIDGE_PDF, output)
    return str(output)

