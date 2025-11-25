from dataclasses import dataclass
from docx import Document


@dataclass
class Paragraph:
    index: int
    text: str
    style: str
    char_count: int
    
    def __repr__(self):
        preview = self.text[:50] + "..." if len(self.text) > 50 else self.text
        return f"[{self.index}|{self.style}|{self.char_count}c] {preview}"


@dataclass
class ResumeStructure:
    paragraphs: list[Paragraph]
    
    def to_prompt_format(self) -> str:
        """Format for LLM prompt - shows structure with constraints."""
        lines = []
        for p in self.paragraphs:
            lines.append(f"[{p.index}|{p.style}|max:{p.char_count}chars] {p.text}")
        return "\n".join(lines)
    
    def get_char_limits(self) -> dict[int, int]:
        """Return mapping of paragraph index to max char count."""
        return {p.index: p.char_count for p in self.paragraphs}


def parse_docx(file_path: str) -> Document:
    """Load and return a DOCX document."""
    return Document(file_path)


def extract_structure(doc: Document) -> ResumeStructure:
    """Extract text and structure from document."""
    paragraphs = []
    
    for i, para in enumerate(doc.paragraphs):
        style_name = para.style.name if para.style else "Normal"
        paragraphs.append(Paragraph(
            index=i,
            text=para.text,
            style=style_name,
            char_count=len(para.text),
        ))
    
    return ResumeStructure(paragraphs=paragraphs)


def print_structure(structure: ResumeStructure):
    """Debug print the structure."""
    print("=== RESUME STRUCTURE ===")
    for p in structure.paragraphs:
        print(p)
    print(f"\nTotal paragraphs: {len(structure.paragraphs)}")
